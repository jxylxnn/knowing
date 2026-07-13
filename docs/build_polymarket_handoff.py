"""Render the Polymarket forecasting handoff Markdown as a polished DOCX."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "polymarket_forecasting_handoff.md"
OUTPUT = ROOT / "polymarket_forecasting_handoff.docx"

INK = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5B6573"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
TABLE_WIDTH = 9360
TABLE_INDENT = 120


def set_run_font(run, *, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_paragraph_shading(paragraph, color):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    p_pr.append(shd)


def set_paragraph_border(paragraph, color="D7DBE2"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_field(paragraph, field_code):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    separate.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(end)
    set_run_font(run, size=9, color=MUTED)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in tokens.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.0

    code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor.from_string(INK)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(4)
    code.paragraph_format.line_spacing = 1.0

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def setup_page(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("POLYMARKET FORECASTING HANDOFF  •  REPOSITORY ARCHITECTURE REVIEW")
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    set_paragraph_border(p)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Prepared 11 July 2026   |   Page ")
    set_run_font(r, size=9, color=MUTED)
    add_field(p, "PAGE")


def add_inline(paragraph, text, *, default_size=11, default_color=INK):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=default_size, color=default_color, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=max(default_size - 1, 8.5), color=DARK_BLUE)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=default_size, color=default_color)


def add_markdown_table(doc, rows):
    parsed = []
    for row in rows:
        cells = [item.strip() for item in row.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        parsed.append(cells)
    if not parsed:
        return
    cols = max(len(row) for row in parsed)
    table = doc.add_table(rows=len(parsed), cols=cols)
    table.style = "Table Grid"
    widths = [TABLE_WIDTH // cols] * cols
    widths[-1] += TABLE_WIDTH - sum(widths)
    set_table_geometry(table, widths)
    for row_index, values in enumerate(parsed):
        tr_pr = table.rows[row_index]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if row_index == 0:
            repeat_header = OxmlElement("w:tblHeader")
            repeat_header.set(qn("w:val"), "true")
            tr_pr.append(repeat_header)
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            if row_index == 0:
                shade(cell, LIGHT_BLUE)
                add_inline(p, value, default_size=9.5, default_color=INK)
                for run in p.runs:
                    run.bold = True
            else:
                add_inline(p, value, default_size=9.5, default_color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render_markdown(doc, lines):
    index = 0
    in_code = False
    while index < len(lines):
        line = lines[index]
        stripped = line.rstrip()
        if stripped.startswith("```"):
            in_code = not in_code
            index += 1
            continue
        if in_code:
            p = doc.add_paragraph(style="Code Block")
            set_paragraph_shading(p, LIGHT_GRAY)
            run = p.add_run(stripped)
            set_run_font(run, name="Consolas", size=8.5, color=INK)
            index += 1
            continue
        if stripped.startswith("|"):
            table_rows = []
            while index < len(lines) and lines[index].rstrip().startswith("|"):
                table_rows.append(lines[index].rstrip())
                index += 1
            add_markdown_table(doc, table_rows)
            continue
        if not stripped:
            index += 1
            continue
        if stripped.startswith("# "):
            index += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
            index += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
            index += 1
            continue
        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.18)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(5)
            set_paragraph_shading(p, CALLOUT)
            add_inline(p, stripped[2:])
            index += 1
            continue
        if re.match(r"^- ", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, stripped[2:])
            index += 1
            continue
        number_match = re.match(r"^(\d+)\. (.*)", stripped)
        if number_match:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, number_match.group(2))
            index += 1
            continue
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].rstrip()
            if (
                not candidate
                or candidate.startswith(("#", "|", "```", "> ", "- "))
                or re.match(r"^\d+\. ", candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(part.strip() for part in paragraph_lines))


def build():
    doc = Document()
    configure_styles(doc)
    setup_page(doc)
    doc.core_properties.title = "Polymarket Forecasting Platform — Architecture Handoff"
    doc.core_properties.subject = "Technical handoff for a point-in-time, domain-pluggable forecasting platform"
    doc.core_properties.author = "Codex"

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(42)
    title.paragraph_format.space_after = Pt(8)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline(title, "Polymarket Forecasting Platform", default_size=26, default_color=INK)
    for run in title.runs:
        run.bold = True
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    add_inline(subtitle, "Architecture handoff — distilled from the NBA prediction repository", default_size=14, default_color=MUTED)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(24)
    add_inline(meta, "Prepared 11 July 2026  •  Read-only-first  •  For the next engineering/model agent", default_size=10, default_color=MUTED)
    lead = doc.add_paragraph()
    lead.paragraph_format.left_indent = Inches(0.18)
    lead.paragraph_format.right_indent = Inches(0.18)
    lead.paragraph_format.space_before = Pt(6)
    lead.paragraph_format.space_after = Pt(8)
    set_paragraph_shading(lead, CALLOUT)
    add_inline(lead, "**Core premise:** Transfer the forecasting discipline, provenance, evaluation, and safety gates—not NBA weights or sports-specific heuristics. Build a calibrated fair-probability engine that is auditable at every decision cutoff.", default_size=11)
    doc.add_page_break()

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    render_markdown(doc, lines)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
